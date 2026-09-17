"""V02-003C installed-only CLI/worker assertions; all data synthetic."""
import sys,json,sqlite3,subprocess,hashlib,os
from pathlib import Path
import scientific_reading
from scientific_reading.library_service import LibraryService
from scientific_reading.models import PaperMetadata
from scientific_reading.xlsx_snapshot import XlsxSnapshotService,XLSX_COLUMNS
from scientific_reading.xlsx_user_fields import BASELINE_SHEET
from scientific_reading.background_models import BackgroundRequest
from scientific_reading.background_store import BackgroundJobStore
from openpyxl import load_workbook
root=Path(sys.argv[1]).resolve();mode=sys.argv[2]
assert (root/'.v02-003c-test-instance').read_text()=='synthetic-only'
assert Path(scientific_reading.__file__).resolve().is_relative_to(root/'releases')
assert not os.environ.get('PYTHONPATH')
cols=['personal_thoughts','understanding_level','user_notes']; names=['个人思考','个人理解程度','用户笔记']
records=[]
def cli(data,*args,code=0):
 p=subprocess.run([sys.executable,'-I','-X','utf8','-m','scientific_reading','--data-root',str(data),*args],capture_output=True,text=True,encoding='utf-8')
 records.append({'args':args,'data':str(data),'code':p.returncode,'stdout':p.stdout,'stderr':p.stderr})
 assert p.returncode==code,records[-1]
 return json.loads(p.stdout)
def read(data):
 with sqlite3.connect(data/'library.sqlite') as c:return dict((r[0],list(r[1:])) for r in c.execute('SELECT paper_id,personal_thoughts,understanding_level,user_notes FROM items ORDER BY paper_id'))
def db(data,pid,values):
 with sqlite3.connect(data/'library.sqlite') as c:c.execute('UPDATE items SET personal_thoughts=?,understanding_level=?,user_notes=? WHERE paper_id=?',(*values,pid))
def edit(data,pid,values):
 target=XlsxSnapshotService(data).target;w=load_workbook(target);s=w['文献'];h={c.value:c.column for c in s[1]}
 row=next(i for i in range(2,s.max_row+1) if s.cell(i,h['文献 ID']).value==pid)
 for n,v in zip(names,values):s.cell(row,h[n]).value=v
 w.save(target);w.close()
def seed(data):
 lib=LibraryService(data);ids=[lib.ingest(PaperMetadata(title='V02-003C synthetic '+str(i),doi='10.9999/v3c-'+str(i)))['paper_id'] for i in range(2)];lib.close()
 for pid in ids:db(data,pid,['B']*3)
 assert cli(data,'xlsx-refresh')['status']=='success'
 return ids
if mode=='cases':
 results=[]
 for label in ['old','excel','conflict0','conflict1','conflict2','legacy-equal','legacy-different']:
  data=root/'f1-cases'/label;ids=seed(data);svc=XlsxSnapshotService(data)
  if label=='old':
   db(data,ids[0],['DB new thought','深入','DB new note']);expect=read(data)
   assert cli(data,'xlsx-refresh')['status']=='success';assert read(data)==expect
  elif label=='excel':
   edit(data,ids[0],['Excel new thought','部分理解','Excel new note'])
   assert cli(data,'xlsx-refresh')['status']=='success';assert read(data)[ids[0]]==['Excel new thought','部分理解','Excel new note']
  elif label.startswith('conflict'):
   field=int(label[-1]);e=['B']*3;d=['B']*3;e[field]='Excel conflict';d[field]='DB conflict'
   edit(data,ids[0],e);edit(data,ids[1],['unrelated row edit','B','B']);db(data,ids[0],d);expect=read(data);before=svc.target.read_bytes()
   for attempt in range(3):
    v=cli(data,'xlsx-refresh');assert v['status']=='pending' and v['updated']==0
    assert {'paper_id':ids[0],'field':cols[field],'code':'user_field_conflict'} in v['conflict_details']
    assert read(data)==expect and svc.target.read_bytes()==before
  else:
   w=load_workbook(svc.target);del w[BASELINE_SHEET];w['_身份']['C1']=None;w['_身份']['D1']=None;w.save(svc.target);w.close()
   if label=='legacy-different':edit(data,ids[0],['legacy edit','B','B'])
   before=svc.target.read_bytes();expect=read(data);v=cli(data,'xlsx-refresh');assert read(data)==expect
   if label=='legacy-different':
    assert v['status']=='pending' and v['conflict_details'][0]['code']=='baseline_missing_difference';assert svc.target.read_bytes()==before
   else:
    assert v['status']=='success';w=load_workbook(svc.target);assert BASELINE_SHEET in w.sheetnames;w.close()
  results.append({'name':label,'passed':True,'database':read(data)})
 print(json.dumps({'engine':scientific_reading.__file__,'executable':sys.executable,'checks':results,'commands':records},ensure_ascii=False))
elif mode=='live-seed':
 data=root/'library';info=json.loads((root/'v02-fixture.json').read_text(encoding='utf-8'));pid=info['paper_id']
 assert cli(data,'xlsx-refresh')['status']=='success'
 lib=LibraryService(data);other=lib.ingest(PaperMetadata(title='V02-003C unrelated row',doi='10.9999/v3c-unrelated'))['paper_id'];lib.close()
 assert cli(data,'xlsx-refresh')['status']=='success'
 vals=read(data)[pid];e=vals.copy();e[2]='Live Excel conflict';d=vals.copy();d[2]='Live DB conflict';edit(data,pid,e);edit(data,other,['must not partially import','','']);db(data,pid,d)
 store=BackgroundJobStore(data);job=store.create_or_get(BackgroundRequest(pid,'xlsx_snapshot','c'*64,{'data_root':str(data)})).job_id
 state={'paper_id':pid,'job_id':job,'database':read(data),'xlsxSha256':hashlib.sha256(XlsxSnapshotService(data).target.read_bytes()).hexdigest()}
 (root/'v3c-conflict.json').write_text(json.dumps(state),encoding='utf-8');print(json.dumps(state))
elif mode=='live-check':
 data=root/'library';saved=json.loads((root/'v3c-conflict.json').read_text());store=BackgroundJobStore(data);job=saved['job_id']
 if store.load_status(job).state=='waiting_user':store.transition(job,'queued')
 p=subprocess.run([sys.executable,'-I','-X','utf8','-m','scientific_reading.worker','--data-root',str(data),'--job-id',job],capture_output=True,text=True,encoding='utf-8');assert p.returncode==2,(p.stdout,p.stderr,p.returncode)
 status=BackgroundJobStore(data).load_status(job).to_dict();assert status['state']=='waiting_user' and status['reason_code']=='xlsx_user_fields_conflict'
 assert status['required_input']['updated']==0
 assert status['required_input']['conflict_details']==[{'paper_id':saved['paper_id'],'field':'user_notes','code':'user_field_conflict'}]
 assert read(data)==saved['database'];assert hashlib.sha256(XlsxSnapshotService(data).target.read_bytes()).hexdigest()==saved['xlsxSha256']
 print(json.dumps({'status':status,'database':read(data),'xlsxSha256':saved['xlsxSha256'],'workerExit':p.returncode},ensure_ascii=False))
