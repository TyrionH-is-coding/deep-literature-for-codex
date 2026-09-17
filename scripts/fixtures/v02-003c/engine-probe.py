"""Synthetic installed-wheel checks; never import engine checkout or call a service.

Parser/translation payload adapted from engine 8e00b334 review fixture formula-outline.
The PDF is valid locally generated text; parser output and translations are fixtures.
"""
import hashlib
import json
import sqlite3
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from scientific_reading.library_service import LibraryService
from scientific_reading.models import PaperMetadata, JobState, StageRecord
from scientific_reading.workspace import PaperWorkspace, atomic_write_json
from scientific_reading.mineru_normalizer import MineruNormalizer
from scientific_reading.full_read_service import FullReadService
from scientific_reading.full_read_models import FULL_TRANSLATION_CONTRACT_VERSION
from scientific_reading.full_read_renderer import FullReadRenderer
from scientific_reading.xlsx_snapshot import XlsxSnapshotService, XLSX_COLUMNS
from scientific_reading.library_backup import backup_library, restore_library

root = Path(sys.argv[1]).resolve()
mode = sys.argv[2]
assert (root / '.v02-003c-test-instance').read_text() == 'synthetic-only'
data = root / 'library'
meta_file = root / 'v02-fixture.json'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def notes():
    with sqlite3.connect(data / 'library.sqlite') as c:
        return list(c.execute('SELECT personal_thoughts,understanding_level,user_notes FROM items WHERE paper_id=?', (info['paper_id'],)).fetchone())
if mode == 'seed':
    payload = json.loads(Path(__file__).with_name('formula-outline.json').read_text(encoding='utf-8'))
    metadata = PaperMetadata.from_dict(payload['metadata'])
    lib = LibraryService(data)
    paper = lib.ingest(metadata)['paper_id']
    lib.close()
    ws = PaperWorkspace.create_for_paper_id(data, paper, metadata)
    # PDF objects, byte offsets and xref are generated locally, without third-party input.
    stream = b'BT /F1 12 Tf 50 780 Td (V02-001 synthetic local PDF) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>', b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>', b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>', b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
    pdf = b'%PDF-1.4\n'; offsets = [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(pdf)); pdf += f'{n} 0 obj\n'.encode()+obj+b'\nendobj\n'
    xref = len(pdf)
    pdf += b'xref\n0 6\n0000000000 65535 f \n'+b''.join(f'{o:010} 00000 n \n'.encode() for o in offsets[1:])+f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    ws.source_pdf.write_bytes(pdf)
    base = ws
    generation = sha(base.source_pdf)[:16]
    ws = PaperWorkspace(root=base.root / 'generations' / generation)
    for directory in [ws.parsed_images, ws.parsed_tables, ws.reading_dir, ws.output_dir]: directory.mkdir(parents=True, exist_ok=True)
    atomic_write_json(ws.metadata_path, metadata.to_dict())
    ws.source_pdf.write_bytes(pdf)
    parsed = ws.parsed_dir / 'mineru'; raw = parsed / 'raw'; raw.mkdir(parents=True, exist_ok=True)
    atomic_write_json(raw / 'fixture_content_list.json', payload['content_items'])
    normalized = MineruNormalizer('v02-local-fixture').normalize(raw, parsed, metadata, sha(ws.source_pdf))
    for name in ['source_map.json', 'parse_report.json']:
        p = parsed / name; value = json.loads(p.read_text()); value['method'] = 'auto'
        if name == 'parse_report.json':
            value['provider'] = 'mineru-local-v1'; value['provider_version'] = 'v02-local-fixture'
        atomic_write_json(p, value)
    stage = StageRecord(status='completed', result={'active_parsed_dir':'parsed/mineru', 'source_sha256':sha(ws.source_pdf), 'method':'auto', 'mineru_version':'v02-local-fixture'})
    ws.save_job(JobState(paper_id=generation, status='parsed', stages={'paper_parse_upgrade': stage}))
    base.save_job(JobState(paper_id=paper, status='parsed', stages={'paper_parse_upgrade': StageRecord(status='completed', result={**stage.result,'active_workspace':'generations/'+generation})}))
    svc = FullReadService(); plan = svc.prepare(ws); translations = {r['block_id']:r for r in payload['translations']}
    for batch in plan.plan['batches']:
        source = json.loads((ws.reading_dir / 'full' / batch['source_file']).read_text(encoding='utf-8'))
        svc.save_translation_batch(ws, {'contract_version':FULL_TRANSLATION_CONTRACT_VERSION, 'batch_id':source['batch_id'], 'source_sha256':source['source_sha256'], 'translations':[dict(block_id=r['block_id'], source_text=r['english'], translation_zh=translations[r['block_id']]['translation_zh'], highlight=translations[r['block_id']]['highlight']) for r in source['blocks']]})
    svc.finalize(ws, payload['review'])
    rendered = FullReadRenderer().render_completed(ws, paper_id=paper)
    reader = Path(rendered['reader_html'])
    lib = LibraryService(data); lib.record_pdf_attachment(paper, sha(ws.source_pdf), len(pdf)); lib.publish_reader(paper, reader.relative_to(base.root).as_posix()); lib.close()
    result = XlsxSnapshotService(data).refresh(); assert result['status']=='success'
    xlsx = Path(result['path']); w=load_workbook(xlsx); s=w['文献']; assert tuple(c.value for c in s[1])==XLSX_COLUMNS
    h={c.value:c.column for c in s[1]}; row=next(i for i in range(2,s.max_row+1) if s.cell(i,h['文献 ID']).value==paper)
    expected=['合成个人思考','部分理解','V02 saved note']
    for col,value in zip(['个人思考','个人理解程度','用户笔记'],expected): s.cell(row,h[col]).value=value
    w.save(xlsx); w.close(); assert XlsxSnapshotService(data).refresh()['status']=='success'
    info={'paper_id':paper, 'expected':expected, 'xlsx':str(xlsx), 'reader':str(reader), 'pdf_sha256':sha(ws.source_pdf), 'reader_sha256':sha(reader), 'normalized_blocks':len(normalized.blocks)}
    meta_file.write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8'); assert notes()==expected
    print(json.dumps(info,ensure_ascii=False))
else:
    info=json.loads(meta_file.read_text(encoding='utf-8'))
    if mode=='read':
        actual=notes(); assert actual==info['expected'], actual
        assert sha(Path(info['reader']))==info['reader_sha256']
        print(json.dumps({'notes':actual,'reader_sha256':info['reader_sha256']},ensure_ascii=False))
    elif mode=='conflict':
        before=notes(); new=['SQLite newer thought','深入','SQLite newer note']
        with sqlite3.connect(data/'library.sqlite') as c: c.execute('UPDATE items SET personal_thoughts=?,understanding_level=?,user_notes=? WHERE paper_id=?',(*new,info['paper_id']))
        assert notes()==new
        result=XlsxSnapshotService(data).refresh(); after=notes()
        print(json.dumps({'old_excel':before,'new_sqlite':new,'refresh':result,'after':after,'newer_preserved':after==new},ensure_ascii=False))
    elif mode=='backup':
        archive=root/'v02-library.zip'; target=root/'restored-library'
        saved=backup_library(data,archive); restored=restore_library(archive,target)
        with zipfile.ZipFile(archive) as z: members=z.namelist()
        # Workspace job absolute paths are intentionally rebased by restore.
        assets=[p for p in (data/'papers').rglob('*') if p.is_file() and not p.name.endswith('.lock') and p.name!='job.json']
        mismatches=[str(p.relative_to(data)) for p in assets if not (target/p.relative_to(data)).is_file() or sha(p)!=sha(target/p.relative_to(data))]
        with sqlite3.connect(target/'library.sqlite') as c: actual=list(c.execute('SELECT personal_thoughts,understanding_level,user_notes FROM items WHERE paper_id=?',(info['paper_id'],)).fetchone())
        assert saved['status']==restored['status']=='completed' and not mismatches and actual==notes()
        print(json.dumps({'backup':saved,'restore':restored,'asset_files':len(assets),'asset_mismatches':mismatches,'notes':actual,'members':members,'native_session_in_archive':any('dsh-home' in n for n in members),'handoff_in_archive':any('handoff' in n for n in members)},ensure_ascii=False))

