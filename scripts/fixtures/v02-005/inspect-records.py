"""Read-only XLSX/SQLite verification. No database mutation."""
import sys,json,sqlite3
from pathlib import Path
from openpyxl import load_workbook
root=Path('C:/tmp/v005/instance');out=Path('C:/tmp/v005/evidence')
if sys.argv[1]=='xlsx-diff':
    before=load_workbook(out/'xlsx-before.xlsx');after=load_workbook(out/'xlsx-edited.xlsx')
    assert before.sheetnames==after.sheetnames
    changes=[]
    for name in before.sheetnames:
        a,b=before[name],after[name]
        assert a.sheet_state==b.sheet_state,(name,'visibility')
        for row in a:
            for cell in row:
                other=b[cell.coordinate]
                if cell.value!=other.value:changes.append({'sheet':name,'cell':cell.coordinate,'before':cell.value,'after':other.value})
    assert len(changes)==3,changes
    assert {c['cell'] for c in changes}=={'G2','H2','I2'}
    print(json.dumps({'only_three_values_changed':True,'changes':changes},ensure_ascii=False))
else:
    with sqlite3.connect(f'file:{(root/"library/library.sqlite").as_posix()}?mode=ro',uri=True) as c:
        row=c.execute('SELECT personal_thoughts,understanding_level,user_notes FROM items WHERE paper_id=?',('title_d8b339a0356f',)).fetchone()
    expected=['合成思考：公式与目录已核对','部分理解','V02-005：重启后继续阅读']
    assert list(row)==expected,row
    print(json.dumps({'personal_records':list(row),'expected':expected,'sqlite_read_only':True},ensure_ascii=False))
